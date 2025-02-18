from pptx import Presentation


def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    full_text = ""

    for slide_number, slide in enumerate(prs.slides, start=1):
        full_text += f"--- Slide {slide_number} ---\n"
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    full_text += text + "\n"
        full_text += "\n"

    return full_text
