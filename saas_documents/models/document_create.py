# -*- coding: utf-8 -*-

import io
import base64
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from odoo import fields, models, api, _


def create_attachment(file_type):
    if file_type == 'docx':
        doc = Document()
        output = io.BytesIO()
        doc.save(output)
    elif file_type == 'xlsx':
        wb = Workbook()
        output = io.BytesIO()
        wb.save(output)
    elif file_type == 'pptx':
        prs = Presentation()
        output = io.BytesIO()
        prs.save(output)
    else:
        return False

    file_data = output.getvalue()
    file_base64 = base64.b64encode(file_data)
    return file_base64


class DocumentsCreate(models.TransientModel):
    _name = 'documents.create'
    _description = 'Documents Create'

    name = fields.Char(string=_('Name'))
    type = fields.Selection(string=_('Type'),
                            selection=[('docx', 'Docs'),
                                       ('xlsx', 'Excel'),
                                       ('pptx', 'PowerPoint'),
                                       ], default='docx')
    folder_id = fields.Many2one('documents.folder', string=_("Workspace"), ondelete="cascade")

    def action_document_create(self):
        self.env['documents.document'].create({
            'name': f'{self.name}.{self.type}',
            'folder_id': self.folder_id.id,
            'attachment_id': self.create_attachment(),
        })

        return {
            'type': 'ir.actions.client',
            'tag': 'reload',
        }

    def create_attachment(self):
        if self.type in ['docx', 'xlsx', 'pptx']:
            attachment_data = create_attachment(self.type)
        else:
            return False

        attachment = self.env['ir.attachment'].create({
            'name': f'{self.name}.{self.type}',
            'datas': attachment_data,
        })
        return attachment.id
